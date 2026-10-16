"""產生 Excel 對帳報表。"""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from reconcile.formats import KIND_LABELS
from reconcile.service import JobView

SHEETS = ("摘要", "已配對", "待確認", "金額不符", "未收款訂單", "不明入帳", "略過的資料列")
RED = PatternFill("solid", fgColor="FDE2E1")
YELLOW = PatternFill("solid", fgColor="FEF3C7")
MONEY = "#,##0"
DATE = "yyyy-mm-dd"


def build_report(view: JobView) -> bytes:
    s = view.summary
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(
        workbook,
        "摘要",
        ["項目", "數值"],
        [
            ("銀行入帳筆數", s["deposit_count"]),
            ("入帳總額", s["deposit_total"]),
            ("訂單筆數", s["order_count"]),
            ("訂單總額", s["order_total"]),
            ("已配對筆數", s["matched_count"]),
            ("配對率（%）", s["match_rate"]),
            ("待確認入帳", s["pending_count"]),
            ("金額不符", s["mismatch_count"]),
            ("未收款筆數", s["unpaid_count"]),
            ("未收款總額", s["unpaid_total"]),
            ("不明入帳筆數", s["unknown_count"]),
            ("不明入帳總額", s["unknown_total"]),
            ("略過的資料列", s["skipped_count"]),
        ],
        money=(2,),
    )
    pair_header = [
        "入帳日期",
        "入帳金額",
        "備註",
        "訂單編號",
        "訂單日期",
        "訂單金額",
        "客戶",
        "依據",
    ]
    _sheet(
        workbook,
        "已配對",
        pair_header,
        [_pair(link) for link in view.matched],
        money=(2, 6),
        dates=(1, 5),
    )
    _sheet(
        workbook,
        "待確認",
        ["入帳列號", *pair_header],
        [(link.deposit.row_no, *_pair(link)) for link in view.pending],
        fill=YELLOW,
        money=(3, 7),
        dates=(2, 6),
    )
    _sheet(
        workbook,
        "金額不符",
        [*pair_header, "差額"],
        [(*_pair(link), link.deposit.amount - link.order.amount) for link in view.mismatched],
        fill=RED,
        money=(2, 6, 9),
        dates=(1, 5),
    )
    _sheet(
        workbook,
        "未收款訂單",
        ["訂單編號", "訂單日期", "訂單金額", "客戶"],
        [(o.order_no, o.date, o.amount, o.customer) for o in view.unpaid],
        fill=RED,
        money=(3,),
        dates=(2,),
    )
    _sheet(
        workbook,
        "不明入帳",
        ["入帳列號", "入帳日期", "入帳金額", "備註", "帳號末五碼"],
        [(d.row_no, d.date, d.amount, d.memo, d.last5 or "") for d in view.unknown],
        fill=RED,
        money=(3,),
        dates=(2,),
    )
    _sheet(
        workbook,
        "略過的資料列",
        ["來源", "原始列號", "原因"],
        [(KIND_LABELS[r.source], r.row_no, r.reason) for r in view.skipped],
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pair(link) -> tuple:
    d, o = link.deposit, link.order
    return (d.date, d.amount, d.memo, o.order_no, o.date, o.amount, o.customer, link.reason)


def _sheet(workbook, title, header, rows, *, fill=None, money=(), dates=()):
    sheet = workbook.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append(list(row))
        for cell in sheet[sheet.max_row]:
            if fill:
                cell.fill = fill
            if cell.column in money:
                cell.number_format = MONEY
            elif cell.column in dates:
                cell.number_format = DATE
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) * 2
        sheet.column_dimensions[column[0].column_letter].width = min(50, max(10, width))
    return sheet
