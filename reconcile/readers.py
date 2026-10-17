"""讀取上傳的 CSV 或 Excel，找出表頭列。"""

import csv
import io
from dataclasses import dataclass

from openpyxl import load_workbook

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000
HEADER_SCAN_ROWS = 10
# 只讀前幾欄；避免工作表的 <dimension> 因遠端格式設定而把每列補成上萬欄
MAX_COLUMNS = 200
# 前幾列中，第一個至少有這麼多個非空白儲存格的列視為表頭（銀行匯出常在上方有標題列）
HEADER_MIN_COLUMNS = 3
# 台灣銀行匯出的 CSV 常是 Big5（cp950）
CSV_ENCODINGS = ("utf-8-sig", "cp950")


class ReadError(Exception):
    """檔案無法讀取；訊息直接顯示給使用者。"""


@dataclass
class Table:
    header: list[str]
    # （原始列號，儲存格）；列號從 1 起算，與 Excel／文字編輯器看到的一致
    rows: list[tuple[int, list]]

    def column(self, row: list, name: str):
        index = self.header.index(name)
        return row[index] if index < len(row) else None


def read_table(filename: str, data: bytes) -> Table:
    if len(data) > MAX_BYTES:
        raise ReadError("檔案超過 5 MB")
    name = filename.lower()
    if name.endswith(".csv"):
        raw = _read_csv(data)
    elif name.endswith((".xlsx", ".xlsm")):
        raw = _read_xlsx(data)
    else:
        raise ReadError("只支援 CSV 或 Excel（.xlsx）檔")
    return _to_table(raw)


def _read_csv(data: bytes) -> list[list]:
    for encoding in CSV_ENCODINGS:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        rows = (_trim([cell.strip() for cell in row]) for row in csv.reader(io.StringIO(text)))
        return _collect(rows)
    raise ReadError("無法判斷 CSV 的文字編碼（支援 UTF-8 與 Big5）")


def _read_xlsx(data: bytes) -> list[list]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ReadError("無法讀取 Excel 檔，檔案可能已損壞") from e
    try:
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        rows = (
            _trim([_clean(value) for value in row])
            for row in sheet.iter_rows(max_col=MAX_COLUMNS, values_only=True)
        )
        return _collect(rows)
    finally:
        workbook.close()


def _collect(rows) -> list[list]:
    # 邊讀邊檢查，超過上限就立刻停止，不把整個檔案讀進記憶體
    limit = HEADER_SCAN_ROWS + MAX_ROWS + 1
    result = []
    for row in rows:
        result.append(row)
        if len(result) > limit:
            raise ReadError("資料超過 2 萬列")
    return result


def _trim(row: list) -> list:
    while row and _is_blank(row[-1]):
        row.pop()
    return row


def _clean(value):
    return value.strip() if isinstance(value, str) else value


def _is_blank(value) -> bool:
    return value is None or value == ""


def _to_table(raw: list[list]) -> Table:
    for index, row in enumerate(raw[:HEADER_SCAN_ROWS]):
        if sum(not _is_blank(value) for value in row) >= HEADER_MIN_COLUMNS:
            header_index = index
            break
    else:
        raise ReadError("找不到表頭列（前 10 列內至少要有一列含 3 個以上的欄名）")
    header = ["" if _is_blank(value) else str(value).strip() for value in raw[header_index]]
    rows = [
        (number + 1, list(row))
        for number, row in enumerate(raw)
        if number > header_index and not all(_is_blank(value) for value in row)
    ]
    if len(rows) > MAX_ROWS:
        raise ReadError("資料超過 2 萬列")
    return Table(header, rows)
