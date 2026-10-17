import io

from openpyxl import load_workbook

from reconcile import service
from reconcile.report import SHEETS, build_report
from tests.test_service import make_job


def test_report_sheets_and_rows(session_factory):
    with session_factory() as session:
        view = service.job_view(make_job(session))
        workbook = load_workbook(io.BytesIO(build_report(view)))
    assert workbook.sheetnames == list(SHEETS)
    assert workbook["摘要"]["A2"].value == "銀行入帳筆數"
    assert workbook["已配對"].max_row == 2  # 標題＋1 筆
    assert workbook["待確認"].max_row == 3  # 標題＋2 筆候選
    assert workbook["未收款訂單"]["A2"].value == "C1001"
    assert workbook["不明入帳"]["C2"].value == 5000
    assert workbook["不明入帳"]["C2"].number_format == "#,##0"
    assert workbook["略過的資料列"]["B2"].value == 5
    assert workbook["已配對"].freeze_panes == "A2"


def _unknown_job(session, memos):
    from reconcile.formats import BANK, ORDERS
    from reconcile.readers import Table
    from tests.test_service import BANK_FMT, ORDER_FMT

    bank = Table(
        ["日期", "存入", "備註", "帳號"],
        [(i + 2, ["2026/10/01", "5000", memo, ""]) for i, memo in enumerate(memos)],
    )
    orders = Table(["編號", "日期", "金額"], [(2, ["Z1", "2026/10/01", "1"])])
    return service.create_job(
        session,
        tables={BANK: bank, ORDERS: orders},
        formats={BANK: BANK_FMT, ORDERS: ORDER_FMT},
        filenames={BANK: "b.csv", ORDERS: "o.xlsx"},
        window_days=7,
    )


def test_report_text_starting_with_equals_is_not_a_formula(session_factory):
    with session_factory() as session:
        view = service.job_view(_unknown_job(session, ["=1+1"]))
        workbook = load_workbook(io.BytesIO(build_report(view)))
    cell = workbook["不明入帳"]["D2"]
    assert cell.value == "=1+1"
    assert cell.data_type == "s"


def test_report_large_sheet_is_fast(session_factory):
    import time

    with session_factory() as session:
        view = service.job_view(_unknown_job(session, ["x"] * 5000))
        start = time.perf_counter()
        data = build_report(view)
        elapsed = time.perf_counter() - start
    assert elapsed < 5
    assert load_workbook(io.BytesIO(data))["不明入帳"].max_row == 5001
